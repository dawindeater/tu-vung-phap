"""Sinh dữ liệu JSON cho web từ file xlsx của gia sư.

CÁCH CHẠY (một lệnh, từ thư mục gốc của dự án):
    py scripts/build_data.py

Đọc : data-src/tu_vung_A1_theo_chu_de_v2.xlsx  và  data-src/confusable.json
Ghi : data/topics.json, data/t01.json ... data/t11.json
File xlsx KHÔNG bao giờ bị sửa. Muốn đổi nội dung: gia sư sửa xlsx rồi chạy lại lệnh trên.
Cần thư viện openpyxl (pip install openpyxl).
"""
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data-src" / "tu_vung_A1_theo_chu_de_v2.xlsx"
CONFUSABLE = ROOT / "data-src" / "confusable.json"
OUT = ROOT / "data"

ARTICLES = ("le ", "la ", "les ", "un ", "une ", "l'")
NOUN_POS = {"nm", "nf", "nm/nf", "nm pl", "nf pl", "np m", "np f", "np m pl"}
NOT_SPLIT = {"Et vous ? / Et toi ?", "être né / née à"}  # dấu "/" ở đây không phải hai giống


def clean(v):
    """NFC, bỏ khoảng trắng thừa. Giữ nguyên chữ gốc (kể cả dấu nháy cong)."""
    if v is None:
        return ""
    s = unicodedata.normalize("NFC", str(v))
    return re.sub(r"\s+", " ", s).strip()


def cmp_key(s):
    """Dạng để so sánh: thường hóa, nháy cong -> thẳng."""
    return clean(s).replace("’", "'").replace("‘", "'").casefold()


def slug(s):
    s = clean(s).casefold().replace("œ", "oe").replace("æ", "ae")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-") or "x"


def norm_pos(pos):
    return "nf pl" if pos == "nfpl" else pos


def pos_group(pos):
    p = norm_pos(pos)
    if p in NOUN_POS:
        return "noun"
    return {"v": "verb", "adj": "adj", "adv": "adv", "expr": "expr"}.get(p, "other")


def gender(pos):
    p = norm_pos(pos)
    return p if p in NOUN_POS else None


def split_forms(fr, pos):
    if " / " in fr and fr not in NOT_SPLIT and norm_pos(pos) in ("nm/nf", "adj"):
        return [clean(x) for x in fr.split(" / ")]
    return [fr]


def strip_article(s):
    k = s.replace("’", "'")
    for a in ARTICLES:
        if k.casefold().startswith(a):
            return s[len(a):].strip()
    return s


def make_accept(forms):
    out = []
    for f in forms:
        for v in (strip_article(f), f):
            v = v.replace("’", "'").replace("‘", "'")
            if v and v not in out:
                out.append(v)
    return out


def is_typeable(forms):
    for f in forms:
        if "…" in f or "..." in f:
            return False
        if f.endswith(("?", "!", ".")):
            return False
    return max(len(f.split()) for f in forms) < 5


def audio_text(forms):
    t = forms[0].replace(" / ", " ")
    t = re.sub(r"\([^)]*\)", " ", t).replace("…", " ").replace("...", " ")
    t = re.sub(r"\s+", " ", t).strip(" ,;")
    return t


def vi_main(vi):
    return clean(vi.split(";")[0]) if ";" in vi else vi


def read_rows():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    topics = []
    for r in wb["Tong_quan"].iter_rows(min_row=2, values_only=True):
        if not r[2]:  # dòng "Tổng"
            break
        topics.append({"num": int(r[0]), "vi": clean(r[1]), "sheet": clean(r[2])})
    for t in topics:
        rows = []
        for r in wb[t["sheet"]].iter_rows(min_row=2, values_only=True):
            if not any(c not in (None, "") for c in r[:5]):
                continue
            rows.append({"group": clean(r[0]), "fr": clean(r[1]), "pos": clean(r[2]),
                         "vi": clean(r[3]), "note": clean(r[4])})
        t["rows"] = rows
    return topics


def make_lessons(rows):
    """Chia các dòng thành bài theo Nhóm con (đúng thứ tự trong xlsx)."""
    groups = []  # [(tên nhóm, [dòng])]
    for r in rows:
        if groups and groups[-1][0] == r["group"]:
            groups[-1][1].append(r)
        else:
            groups.append((r["group"], [r]))
    lessons = []  # [{"names": [...], "rows": [...]}]
    i = 0
    while i < len(groups):
        name, rs = groups[i]
        n = len(rs)
        if n > 18:
            k = math.ceil(n / 15)
            base, extra = divmod(n, k)
            pos = 0
            for j in range(k):
                size = base + (1 if j < extra else 0)
                lessons.append({"names": [f"{name} ({j + 1}/{k})"], "rows": rs[pos:pos + size]})
                pos += size
            i += 1
        elif n >= 10:
            lessons.append({"names": [name], "rows": list(rs)})
            i += 1
        else:
            names, acc = [name], list(rs)
            i += 1
            while len(acc) < 10 and i < len(groups) and len(acc) + len(groups[i][1]) <= 18:
                names.append(groups[i][0])
                acc += groups[i][1]
                i += 1
            lessons.append({"names": names, "rows": acc})
    # Bài dưới 6 từ: dồn vào bài liền trước (hoặc liền sau nếu là bài đầu) khi tổng không quá 22.
    # (Quy tắc gốc chỉ nói bài cuối; mở rộng cho bài giữa vì có trường hợp nhóm nhỏ kẹt trước nhóm lớn.)
    i = 0
    while i < len(lessons):
        cur = lessons[i]
        if len(cur["rows"]) < 6 and len(lessons) >= 2:
            if i > 0 and len(lessons[i - 1]["rows"]) + len(cur["rows"]) <= 22:
                lessons[i - 1]["names"] += cur["names"]
                lessons[i - 1]["rows"] += cur["rows"]
                lessons.pop(i)
                continue
            if i + 1 < len(lessons) and len(lessons[i + 1]["rows"]) + len(cur["rows"]) <= 22:
                lessons[i + 1]["names"] = cur["names"] + lessons[i + 1]["names"]
                lessons[i + 1]["rows"] = cur["rows"] + lessons[i + 1]["rows"]
                lessons.pop(i)
                continue
        i += 1
    return lessons


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    topics = read_rows()

    conf = {}
    conf_unmatched = []
    if CONFUSABLE.exists():
        data = json.loads(CONFUSABLE.read_text(encoding="utf-8"))
        for g in data["groups"]:
            for it in g["items"]:
                conf.setdefault(cmp_key(it), []).append(g["id"])
    conf_hit = set()

    # sameMeaning: đếm viMain trên toàn bộ dữ liệu
    vm_count = {}
    for t in topics:
        for r in t["rows"]:
            k = cmp_key(vm := vi_main(r["vi"]))
            vm_count[k] = vm_count.get(k, 0) + 1

    OUT.mkdir(exist_ok=True)
    summary = []
    for t in topics:
        tcode = f"t{t['num']:02d}"
        seen = {}
        lessons_out = []
        for li, les in enumerate(make_lessons(t["rows"]), 1):
            words = []
            for r in les["rows"]:
                forms = split_forms(r["fr"], r["pos"])
                base = f"{tcode}-{slug(r['fr'])}"
                seen[base] = seen.get(base, 0) + 1
                wid = base if seen[base] == 1 else f"{base}-{seen[base]}"
                vm = vi_main(r["vi"])
                same = f"sm-{slug(vm)}" if vm_count[cmp_key(vm)] > 1 else None
                keys = [cmp_key(r["fr"])] + [cmp_key(f) for f in forms]
                cg = None
                for k in keys:
                    if k in conf:
                        cg = conf[k][0]
                        conf_hit.add(k)
                        break
                words.append({
                    "id": wid, "fr": r["fr"], "forms": forms, "pos": r["pos"],
                    "posGroup": pos_group(r["pos"]), "gender": gender(r["pos"]),
                    "vi": r["vi"], "viMain": vm, "note": r["note"],
                    "accept": make_accept(forms), "typeable": is_typeable(forms),
                    "audioText": audio_text(forms),
                    "sameMeaningGroup": same, "confusableGroup": cg,
                })
            lessons_out.append({"id": f"{tcode}-b{li}", "title": " · ".join(les["names"]), "words": words})
        topic_json = {"id": t["num"], "vi": t["vi"], "frName": "", "lessons": lessons_out}
        (OUT / f"{tcode}.json").write_text(
            json.dumps(topic_json, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        summary.append({
            "id": t["num"], "vi": t["vi"], "frName": "",
            "lessonCount": len(lessons_out),
            "wordCount": sum(len(l["words"]) for l in lessons_out),
            "lessons": [{"id": l["id"], "title": l["title"], "wordCount": len(l["words"])} for l in lessons_out],
        })
    (OUT / "topics.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    for k in conf:
        if k not in conf_hit:
            print(f"CẢNH BÁO: mục trong confusable.json không có trong xlsx: {k}")
    total_w = sum(s["wordCount"] for s in summary)
    total_l = sum(s["lessonCount"] for s in summary)
    print(f"Xong: {len(summary)} chủ điểm, {total_l} bài, {total_w} từ -> thư mục data/")


if __name__ == "__main__":
    main()
